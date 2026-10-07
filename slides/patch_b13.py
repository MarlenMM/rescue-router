"""Batch 13.6: update the numbers in Marlen's exported deck (Rescue Router.pptx, 7 Oct 16:06) to the final model.

    python slides/patch_b13.py --minutes 35 --upload_s 20      # -> slides/20260876_MarlenMelis_slides.pptx

Text is replaced inside the existing runs (formatting kept); a phrase split across runs is rewritten into the
paragraph's first run. Slide text and speaker notes; the curve and the 7301 / 7241 scene images are swapped for the
regenerated ones (same size and format). The source file is not modified.
"""
import argparse
import io

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# (slide number, picture index on that slide) -> regenerated asset (slides/make_slide_assets.py, Batch 13 model)
IMAGES = {(5, 0): 'slides/assets/curve.png', (6, 0): 'slides/assets/route_7301_pred.jpg',
          (6, 1): 'slides/assets/route_7301_gt.jpg', (7, 0): 'slides/assets/fail_7241_pred.jpg'}
ALT = {(5, 0): 'Learning curve: mIoU on 100 test photos rises from 0.26 with 1 labelled photo to 0.52 with 50 (final model, '
               'tuned decode), with 95 percent intervals; full test 0.507 (first model 0.472); PSPNet with 1,445 labels at 0.848'}


def swap_images(prs):
    n = 0
    for k, slide in enumerate(prs.slides, 1):
        pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
        for j, pic in enumerate(pics):
            if (k, j) not in IMAGES:
                continue
            part = pic.part.related_part(pic._element.blip_rId)
            im = Image.open(IMAGES[(k, j)]).convert('RGB')
            buf = io.BytesIO()
            im.save(buf, 'PNG' if part.content_type == 'image/png' else 'JPEG', quality=90)
            part._blob = buf.getvalue()
            if (k, j) in ALT:
                pic._element.xpath('./p:nvPicPr/p:cNvPr')[0].set('descr', ALT[(k, j)])
            n += 1
    return n


def pairs(a):
    return [
        # slide 2 (1/6): the three headline numbers
        ('.472 mIoU', '.507 mIoU'),
        ('.759 F1', '.752 F1'),
        # slide 4 (3/6)
        ('Fine-tuned: biases, matching temperature, one LayerNorm: 5 MB per model. No other pretrained model.',
         'Fine-tuned: biases, temperature, one LayerNorm (5 MB per model). Plus SAM to redraw building outlines.'),
        ('7 runs on free Kaggle T4 GPUs, about 17 GPU-hours.', '28 runs on free Kaggle T4 GPUs, about 29 GPU-hours.'),
        ('reruns everything on a fresh T4 in about 30 minutes;',
         f'reruns everything on a fresh T4 in about {a.minutes} minutes;'),
        # slide 5 (4/6)
        ('Flooded roads are the gap: IoU .29 vs .82 for PSPNet.', 'Flooded roads are the gap: IoU .31 vs .82 for PSPNet.'),
        # slide 6 (5/6)
        ('F1 .759 on 1,232 test houses (.835 on the 100-photo set, .818 RescueNet)',
         'F1 .752 on 1,232 test houses (.819 on the 100-photo set, .818 RescueNet)'),
        ('“confident: truck” is right 98.7 %, “check first” only 74 %', '“confident: truck” is right 100 %, “check first” 79 %'),
        ('runs the model on an upload in 16 s', f'runs the model on an upload in {a.upload_s} s'),
        # slide 7 (6/6)
        ('Tried and dropped:', 'Tried 36 ideas:'),               # label run and body run of the same paragraph
        (' test-time flips (−.017 with their own tuning), 10 instead of 5 labels (no gain), smoothing.',
         ' kept a shorter schedule, weight averaging, an edge-aware decode and SAM (.472 → .507); the rest lost.'),
        # speaker notes
        ('.472 mIoU on the whole FloodNet test set, a router that finds cut-off houses with F1 .759, and when the system says '
         '"confident: truck", it is right 98.7 percent of the time.',
         '.507 mIoU on the whole FloodNet test set, a router that finds cut-off houses with F1 .752, and when the system says '
         '"confident: truck", it was right for all 208 such houses in the test photos.'),
        ('The only pretrained weights are the course\'s Chameleon checkpoint;',
         'The main pretrained weights are the course\'s Chameleon checkpoint, plus SAM to redraw building outlines;'),
        ('Seven runs, about 17 GPU-hours on Kaggle. And the final notebook reruns on a fresh T4 in about half an hour.',
         f'Seven runs plus 21 improvement runs, about 29 GPU-hours on Kaggle. And the final notebook reruns on a fresh T4 in '
         f'about {a.minutes} minutes.'),
        ('On the full test set the 50-shot model reaches .472.',
         'On the full test set the final 50-shot model reaches .507, up from .472 for my first model after a round of 36 '
         'tested improvements.'),
        ('Most of the gap is flooded roads: .29 against .82', 'Most of the gap is flooded roads: .31 against .82'),
        ('with F1 .759 on the full test.', 'with F1 .752 on the full test; the better map did not change routing measurably.'),
        ('houses marked "confident: truck" are right 98.7 percent of the time; "check first" only 74 percent.',
         'houses marked "confident: truck" were all right; "check first" 79 percent.'),
        ('runs the model live on an uploaded photo in 16 seconds.', f'runs the model live on an uploaded photo in {a.upload_s} seconds.'),
        ('Things I tried and dropped: test-time flips, which looked good on the tuning photos but lost on the test set; more '
         'labels between 5 and 10, which gave nothing measurable; and smoothing.',
         'I also tried 36 improvement ideas on the 50-shot model, each judged only on the 30 tuning photos. Four survived: a '
         'shorter training schedule, averaging two snapshots, an edge-aware decode, and SAM to redraw building outlines; '
         'together .472 to .507 on the test set. Most others lost, for example pseudo-labels and rare-class losses, and '
         'test-time flips looked good on the tuning photos but lost on the test set.'),
        ('98.7 %', '100 %'),                       # slide 2 (last: the slide-6 sentence above contains it too)
    ]


def replace_in_paragraph(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    full = ''.join(r.text for r in par.runs)
    if old in full and par.runs:
        par.runs[0].text = full.replace(old, new)
        for r in par.runs[1:]:
            r.text = ''
        return True
    return False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default='Rescue Router.pptx')
    ap.add_argument('--out', default='slides/20260876_MarlenMelis_slides.pptx')
    ap.add_argument('--minutes', required=True)
    ap.add_argument('--upload_s', required=True)
    a = ap.parse_args()
    prs = Presentation(a.src)
    todo = sorted(pairs(a), key=lambda kv: -len(kv[0]))    # longest first: short patterns also occur inside long ones
    done = set()
    for s in prs.slides:
        frames = [sh.text_frame for sh in s.shapes if sh.has_text_frame]
        if s.has_notes_slide:
            frames.append(s.notes_slide.notes_text_frame)
        for tf in frames:
            for par in tf.paragraphs:
                for k, (old, new) in enumerate(todo):
                    if replace_in_paragraph(par, old, new):
                        done.add(k)
    missing = [todo[k][0][:60] for k in range(len(todo)) if k not in done]
    n_img = swap_images(prs)
    prs.save(a.out)
    print(f'{len(done)}/{len(todo)} replacements; missing: {missing}; {n_img} images swapped; wrote {a.out}')


if __name__ == '__main__':
    main()
